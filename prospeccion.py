import streamlit as st
import sqlite3
import pandas as pd
import plotly.express as px
import json
import hashlib
import re
import io
import smtplib
import time
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
from datetime import datetime
from zoneinfo import ZoneInfo

st.set_page_config(page_title="Prospección en Frío", page_icon="📧", layout="wide")

ZONA_CDMX = ZoneInfo("America/Mexico_City")
DB_PATH = "prospeccion.db"

COLUMNAS_DEFAULT = ["Nombre", "Empresa", "Correo"]
ASUNTO_DEFAULT = "Oportunidad de colaboración con (Empresa)"
CUERPO_DEFAULT = (
    "Hola (Nombre),\n\n"
    "Espero que te encuentres muy bien.\n\n"
    "Me pongo en contacto contigo de parte de [Tu empresa].\n\n"
    "[Escribe aquí tu propuesta de valor...]\n\n"
    "Quedamos a tus órdenes,\n"
    "[Tu nombre y cargo]"
)
PAGINAS = ["Inicio", "Envío masivo", "Envío individual", "Dashboard",
           "Seguimientos", "Mi configuración", "Administración"]


def hoy_cdmx():
    return datetime.now(ZONA_CDMX).date()


def ahora_cdmx():
    return datetime.now(ZONA_CDMX).isoformat()


def hash_pw(p):
    return hashlib.sha256(p.encode()).hexdigest()


# ─────────────────────────────────────────────
# BASE DE DATOS
# ─────────────────────────────────────────────

def get_conn():
    return sqlite3.connect(DB_PATH, check_same_thread=False)


def _add_col(conn, table, col, coltype):
    c = conn.cursor()
    c.execute(f"PRAGMA table_info({table})")
    if col not in [r[1] for r in c.fetchall()]:
        c.execute(f"ALTER TABLE {table} ADD COLUMN {col} {coltype}")
        conn.commit()


def init_db():
    conn = get_conn()
    c = conn.cursor()

    c.execute("""CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        nombre TEXT,
        activo INTEGER DEFAULT 1,
        creado_en TEXT NOT NULL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS configuracion (
        clave TEXT PRIMARY KEY,
        valor TEXT NOT NULL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS campanas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL,
        asunto_plantilla TEXT NOT NULL,
        cuerpo_plantilla TEXT NOT NULL,
        adjunto_nombre TEXT,
        tipo TEXT NOT NULL,
        usuario TEXT NOT NULL,
        creado_en TEXT NOT NULL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS correos_enviados (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        campana_id INTEGER NOT NULL,
        fila_datos TEXT,
        nombres TEXT NOT NULL,
        correos TEXT NOT NULL,
        asunto_final TEXT NOT NULL,
        cuerpo_final TEXT NOT NULL,
        estado TEXT NOT NULL,
        error TEXT,
        enviado_en TEXT NOT NULL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS seguimientos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        correo_enviado_id INTEGER NOT NULL,
        asunto TEXT NOT NULL,
        cuerpo TEXT NOT NULL,
        estado TEXT NOT NULL,
        error TEXT,
        enviado_en TEXT NOT NULL
    )""")

    conn.commit()

    # Migraciones: columnas por usuario
    for col, tipo in [
        ("email_remitente", "TEXT DEFAULT ''"),
        ("email_password", "TEXT DEFAULT ''"),
        ("asunto_plantilla", "TEXT DEFAULT ''"),
        ("cuerpo_plantilla", "TEXT DEFAULT ''"),
    ]:
        _add_col(conn, "usuarios", col, tipo)

    # Usuario admin por defecto
    c.execute("SELECT COUNT(*) FROM usuarios")
    if c.fetchone()[0] == 0:
        c.execute(
            "INSERT INTO usuarios (username, password_hash, nombre, creado_en) VALUES (?,?,?,?)",
            ("admin", hash_pw("admin123"), "Administrador", ahora_cdmx())
        )

    # Config global por defecto
    defaults = {
        "columnas_excel": json.dumps(COLUMNAS_DEFAULT),
        "columna_correo": "Correo",
        "admin_password": hash_pw("admin123"),
        # Credenciales globales de respaldo (opcionales)
        "email_remitente": "",
        "email_password": "",
    }
    for k, v in defaults.items():
        c.execute("INSERT OR IGNORE INTO configuracion (clave, valor) VALUES (?,?)", (k, v))

    # Migración: columna respondido para bases de datos existentes
    c.execute("PRAGMA table_info(correos_enviados)")
    cols = [r[1] for r in c.fetchall()]
    if "respondido" not in cols:
        c.execute("ALTER TABLE correos_enviados ADD COLUMN respondido INTEGER DEFAULT 0")

    conn.commit()
    conn.close()


def get_config(clave, default=None):
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT valor FROM configuracion WHERE clave=?", (clave,))
    row = c.fetchone()
    conn.close()
    return row[0] if row else default


def set_config(clave, valor):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "INSERT INTO configuracion (clave, valor) VALUES (?,?) "
        "ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor",
        (clave, str(valor))
    )
    conn.commit()
    conn.close()


def get_columnas_excel():
    raw = get_config("columnas_excel")
    if not raw:
        return list(COLUMNAS_DEFAULT)
    try:
        return json.loads(raw)
    except Exception:
        return list(COLUMNAS_DEFAULT)


# ── Datos por usuario ──────────────────────────

def get_usuario_row(username):
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT * FROM usuarios WHERE username=?", (username,))
    row = c.fetchone()
    cols = [d[0] for d in c.description] if c.description else []
    conn.close()
    if row and cols:
        return dict(zip(cols, row))
    return {}


def set_usuario_email(username, remitente, password):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "UPDATE usuarios SET email_remitente=?, email_password=? WHERE username=?",
        (remitente.strip(), password.strip(), username)
    )
    conn.commit()
    conn.close()


def set_usuario_plantilla(username, asunto, cuerpo):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "UPDATE usuarios SET asunto_plantilla=?, cuerpo_plantilla=? WHERE username=?",
        (asunto, cuerpo, username)
    )
    conn.commit()
    conn.close()


def get_email_credentials_usuario(username):
    """Credenciales del usuario. Si no tiene, usa las globales o secrets."""
    row = get_usuario_row(username)
    rem = row.get("email_remitente", "").strip()
    pwd = row.get("email_password", "").strip()
    if rem and pwd:
        return rem, pwd
    # Fallback: config global
    rem_g = get_config("email_remitente", "").strip()
    pwd_g = get_config("email_password", "").strip()
    if rem_g and pwd_g:
        return rem_g, pwd_g
    # Fallback: secrets
    try:
        return st.secrets["email"]["remitente"], st.secrets["email"]["password"]
    except Exception:
        return "", ""


def get_plantilla_usuario(username):
    """Plantilla del usuario. Si no tiene, usa los valores por defecto globales."""
    row = get_usuario_row(username)
    asunto = (row.get("asunto_plantilla") or "").strip()
    cuerpo = (row.get("cuerpo_plantilla") or "").strip()
    return (
        asunto if asunto else ASUNTO_DEFAULT,
        cuerpo if cuerpo else CUERPO_DEFAULT,
    )


# ── Usuarios CRUD ──────────────────────────────

def verificar_usuario(username, password):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "SELECT id, nombre FROM usuarios WHERE username=? AND password_hash=? AND activo=1",
        (username.strip(), hash_pw(password))
    )
    row = c.fetchone()
    conn.close()
    return row


def get_usuarios():
    conn = get_conn()
    df = pd.read_sql_query(
        "SELECT id, username, nombre, activo, email_remitente, creado_en FROM usuarios ORDER BY username",
        conn
    )
    conn.close()
    return df


def crear_usuario(username, password, nombre):
    conn = get_conn()
    c = conn.cursor()
    try:
        c.execute(
            "INSERT INTO usuarios (username, password_hash, nombre, creado_en) VALUES (?,?,?,?)",
            (username.strip(), hash_pw(password), nombre.strip(), ahora_cdmx())
        )
        conn.commit()
        conn.close()
        return True, "Usuario creado."
    except sqlite3.IntegrityError:
        conn.close()
        return False, "Ya existe un usuario con ese nombre."


def cambiar_password_usuario(user_id, nueva_password):
    conn = get_conn()
    c = conn.cursor()
    c.execute("UPDATE usuarios SET password_hash=? WHERE id=?", (hash_pw(nueva_password), user_id))
    conn.commit()
    conn.close()


def toggle_usuario_activo(user_id, activo):
    conn = get_conn()
    c = conn.cursor()
    c.execute("UPDATE usuarios SET activo=? WHERE id=?", (int(activo), user_id))
    conn.commit()
    conn.close()


# ── Campañas ───────────────────────────────────

def crear_campana(nombre, asunto, cuerpo, tipo, usuario, adjunto_nombre=None):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "INSERT INTO campanas (nombre, asunto_plantilla, cuerpo_plantilla, adjunto_nombre, tipo, usuario, creado_en) "
        "VALUES (?,?,?,?,?,?,?)",
        (nombre, asunto, cuerpo, adjunto_nombre, tipo, usuario, ahora_cdmx())
    )
    campana_id = c.lastrowid
    conn.commit()
    conn.close()
    return campana_id


def registrar_correo(campana_id, nombres, correos, asunto_final, cuerpo_final, estado, fila_datos=None, error=None):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "INSERT INTO correos_enviados "
        "(campana_id, fila_datos, nombres, correos, asunto_final, cuerpo_final, estado, error, enviado_en) "
        "VALUES (?,?,?,?,?,?,?,?,?)",
        (
            campana_id,
            json.dumps(fila_datos, ensure_ascii=False) if fila_datos else None,
            nombres, correos, asunto_final, cuerpo_final, estado, error, ahora_cdmx()
        )
    )
    conn.commit()
    conn.close()


def registrar_seguimiento(correo_enviado_id, asunto, cuerpo, estado, error=None):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "INSERT INTO seguimientos (correo_enviado_id, asunto, cuerpo, estado, error, enviado_en) "
        "VALUES (?,?,?,?,?,?)",
        (correo_enviado_id, asunto, cuerpo, estado, error, ahora_cdmx())
    )
    conn.commit()
    conn.close()


def get_campanas(solo_usuario=None):
    """Si solo_usuario=None trae todas; si es un username filtra por ese usuario."""
    conn = get_conn()
    where = "WHERE c.usuario=?" if solo_usuario else ""
    params = (solo_usuario,) if solo_usuario else ()
    df = pd.read_sql_query(
        f"""SELECT c.id, c.nombre, c.tipo, c.usuario, c.creado_en,
            COUNT(ce.id) AS total,
            SUM(CASE WHEN ce.estado='enviado' THEN 1 ELSE 0 END) AS exitosos,
            SUM(CASE WHEN ce.estado='error' THEN 1 ELSE 0 END) AS errores
           FROM campanas c
           LEFT JOIN correos_enviados ce ON c.id = ce.campana_id
           {where}
           GROUP BY c.id ORDER BY c.creado_en DESC""",
        conn, params=params
    )
    conn.close()
    return df


def get_correos_campana(campana_id, solo_exitosos=False):
    conn = get_conn()
    q = "SELECT * FROM correos_enviados WHERE campana_id=?"
    if solo_exitosos:
        q += " AND estado='enviado'"
    q += " ORDER BY enviado_en"
    df = pd.read_sql_query(q, conn, params=(campana_id,))
    conn.close()
    return df


def marcar_respondido(correo_id, respondido):
    conn = get_conn()
    c = conn.cursor()
    c.execute("UPDATE correos_enviados SET respondido=? WHERE id=?", (int(respondido), correo_id))
    conn.commit()
    conn.close()


def get_seguimientos_por_correo(correo_enviado_id):
    conn = get_conn()
    df = pd.read_sql_query(
        "SELECT * FROM seguimientos WHERE correo_enviado_id=? ORDER BY enviado_en",
        conn, params=(correo_enviado_id,)
    )
    conn.close()
    return df


def borrar_campana(campana_id):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "DELETE FROM seguimientos WHERE correo_enviado_id IN "
        "(SELECT id FROM correos_enviados WHERE campana_id=?)", (campana_id,)
    )
    c.execute("DELETE FROM correos_enviados WHERE campana_id=?", (campana_id,))
    c.execute("DELETE FROM campanas WHERE id=?", (campana_id,))
    conn.commit()
    conn.close()


def borrar_correo_enviado(correo_id):
    conn = get_conn()
    c = conn.cursor()
    c.execute("DELETE FROM seguimientos WHERE correo_enviado_id=?", (correo_id,))
    c.execute("DELETE FROM correos_enviados WHERE id=?", (correo_id,))
    conn.commit()
    conn.close()


def leer_respaldo_db():
    with open(DB_PATH, "rb") as f:
        return f.read()


# ─────────────────────────────────────────────
# CORREO Y PLACEHOLDERS
# ─────────────────────────────────────────────

def formatear_multi(valor_str):
    partes = [p.strip() for p in str(valor_str).split(';') if p.strip()]
    if not partes:
        return str(valor_str)
    if len(partes) == 1:
        return partes[0]
    if len(partes) == 2:
        return f"{partes[0]} y {partes[1]}"
    return ", ".join(partes[:-1]) + f" y {partes[-1]}"


def aplicar_placeholders(texto, fila_dict, columnas):
    for col in columnas:
        valor_raw = str(fila_dict.get(col, ''))
        valor = formatear_multi(valor_raw) if ';' in valor_raw else valor_raw
        texto = texto.replace(f"({col})", valor)
    return texto


def validar_placeholders(asunto, cuerpo, columnas):
    encontrados = set(re.findall(r'\(([^)]+)\)', asunto + cuerpo))
    return [p for p in encontrados if p not in columnas]


def generar_plantilla_excel(columnas):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospectos"
    ws.append(list(columnas))
    ejemplo = []
    for col in columnas:
        col_lower = str(col).lower()
        if any(x in col_lower for x in ["correo", "email", "mail"]):
            ejemplo.append("ejemplo@empresa.com")
        elif any(x in col_lower for x in ["nombre", "name"]):
            ejemplo.append("Juan Pérez")
        elif any(x in col_lower for x in ["empresa", "company"]):
            ejemplo.append("Empresa Ejemplo")
        else:
            ejemplo.append(f"Ejemplo {col}")
    ws.append(ejemplo)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def enviar_correo(destinatarios, asunto, cuerpo, adjuntos=None, username=None):
    remitente, password = get_email_credentials_usuario(username or "")
    if not remitente or not password:
        return False, "Sin credenciales de correo. Configúralas en Mi configuración → Correo."
    try:
        msg = MIMEMultipart()
        msg['From'] = remitente
        msg['To'] = ', '.join(destinatarios)
        msg['Subject'] = asunto
        msg.attach(MIMEText(cuerpo, 'plain', 'utf-8'))
        if adjuntos:
            for adj_bytes, adj_nombre in adjuntos:
                if not adj_bytes or not adj_nombre:
                    continue
                part = MIMEApplication(adj_bytes, Name=adj_nombre)
                part['Content-Disposition'] = f'attachment; filename="{adj_nombre}"'
                msg.attach(part)
        with smtplib.SMTP_SSL('smtp.gmail.com', 465) as srv:
            srv.login(remitente, password)
            srv.sendmail(remitente, destinatarios, msg.as_string())
        return True, "Enviado."
    except Exception as e:
        return False, str(e)


# ─────────────────────────────────────────────
# AUTH
# ─────────────────────────────────────────────

def login_form():
    col_l, col_c, col_r = st.columns([1, 1.2, 1])
    with col_c:
        st.markdown("<br><br><br>", unsafe_allow_html=True)
        st.markdown("""
        <div style="text-align:center; margin-bottom:2.5rem;">
            <div style="
                display:inline-block;
                background:#111111; color:#ff3c00;
                width:52px; height:52px; line-height:52px;
                font-size:24px; margin-bottom:1.2rem;
                font-weight:700;
            ">✉</div>
            <h1 style="
                font-family:'Space Grotesk',sans-serif;
                font-size:2rem; font-weight:700; text-transform:uppercase;
                color:#111111; margin:0 0 0.4rem 0;
                letter-spacing:-0.04em; border:none; padding:0;
            ">Prospección en Frío</h1>
            <p style="font-family:'Space Mono',monospace; color:#999999;
                      font-size:0.7rem; margin:0; letter-spacing:0.1em; text-transform:uppercase;">
                Ingresa tus credenciales
            </p>
        </div>
        """, unsafe_allow_html=True)
        with st.form("form_login"):
            username = st.text_input("Usuario")
            password = st.text_input("Contraseña", type="password")
            st.markdown("<br>", unsafe_allow_html=True)
            if st.form_submit_button("Entrar →", type="primary"):
                resultado = verificar_usuario(username, password)
                if resultado:
                    st.session_state.logged_in = True
                    st.session_state.username = username.strip()
                    st.session_state.user_nombre = resultado[1] or username
                    st.session_state.pagina = "Inicio"
                    st.rerun()
                else:
                    st.error("Usuario o contraseña incorrectos.")


def verificar_admin():
    if "is_admin" not in st.session_state:
        st.session_state.is_admin = False
    if not st.session_state.is_admin:
        st.subheader("🔒 Acceso de administrador")
        pwd = st.text_input("Contraseña de administrador", type="password", key="admin_pwd_input")
        if st.button("Entrar como administrador"):
            if hash_pw(pwd) == get_config("admin_password", hash_pw("admin123")):
                st.session_state.is_admin = True
                st.rerun()
            else:
                st.error("Contraseña incorrecta.")
        st.stop()


# ─────────────────────────────────────────────
# INIT
# ─────────────────────────────────────────────

init_db()

# ─── Tema visual Cultural / Experimental ─────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@300;400;500;600;700&family=Space+Mono:wght@400;700&display=swap');

html, body, [class*="css"] {
    font-family: 'Space Grotesk', sans-serif;
    font-weight: 400;
}

/* ── Fondo: blanco roto con textura ── */
.stApp {
    background-color: #f2f0eb;
    color: #111111;
}

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background-color: #111111 !important;
    border-right: none !important;
}
[data-testid="stSidebar"] p,
[data-testid="stSidebar"] span,
[data-testid="stSidebar"] label,
[data-testid="stSidebar"] div:not([data-testid]),
[data-testid="stSidebar"] a { color: #f2f0eb !important; }

/* ── Texto principal (evita que herede color blanco del sidebar) ── */
.main .block-container { color: #111111 !important; }
.main p, .main span:not([class*="css"]), .main li { color: #111111 !important; }
[data-testid="stTabsContent"] * { color: #111111 !important; }
[data-testid="stMarkdownContainer"] { color: #111111 !important; }
[data-testid="stVerticalBlock"] { color: #111111 !important; }
[data-testid="stSidebar"] .stRadio label {
    font-size: 0.82rem !important;
    font-weight: 500 !important;
    letter-spacing: 0.04em !important;
    color: #aaaaaa !important;
}

/* ── Títulos ── */
h1 {
    font-family: 'Space Grotesk', sans-serif !important;
    font-weight: 700 !important;
    font-size: 2.4rem !important;
    color: #111111 !important;
    letter-spacing: -0.04em;
    border-bottom: 3px solid #111111;
    padding-bottom: 0.6rem;
    margin-bottom: 1.4rem !important;
    text-transform: uppercase;
}
h2 {
    font-weight: 700 !important;
    color: #111111 !important;
    text-transform: uppercase;
    letter-spacing: -0.02em;
}
h3 {
    font-family: 'Space Mono', monospace !important;
    font-size: 0.75rem !important;
    font-weight: 400 !important;
    color: #888888 !important;
    letter-spacing: 0.08em !important;
    text-transform: uppercase !important;
}

/* ── Botón primario: negro total, hover con acento ── */
.stButton > button[kind="primary"] {
    background: #111111 !important;
    color: #f2f0eb !important;
    border: 2px solid #111111 !important;
    border-radius: 0px !important;
    font-weight: 700 !important;
    font-size: 0.78rem !important;
    letter-spacing: 0.12em !important;
    text-transform: uppercase !important;
    padding: 0.55rem 1.6rem !important;
    transition: all 0.15s ease !important;
}
.stButton > button[kind="primary"]:hover {
    background: #ff3c00 !important;
    border-color: #ff3c00 !important;
    color: #ffffff !important;
    transform: none !important;
}

/* ── Botón secundario ── */
.stButton > button[kind="secondary"] {
    background: transparent !important;
    color: #555555 !important;
    border: 2px solid #cccccc !important;
    border-radius: 0px !important;
    font-size: 0.78rem !important;
    font-weight: 600 !important;
    letter-spacing: 0.08em !important;
    text-transform: uppercase !important;
    transition: all 0.15s ease !important;
}
.stButton > button[kind="secondary"]:hover {
    border-color: #111111 !important;
    color: #111111 !important;
}

/* ── Inputs ── */
.stTextInput > div > div > input,
.stTextArea > div > div > textarea,
.stSelectbox > div > div > div,
.stNumberInput > div > div > input {
    background: #ffffff !important;
    border: 2px solid #cccccc !important;
    border-radius: 0px !important;
    color: #111111 !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-size: 0.9rem !important;
    transition: border-color 0.15s ease !important;
}
.stTextInput > div > div > input:focus,
.stTextArea > div > div > textarea:focus {
    border-color: #ff3c00 !important;
    box-shadow: none !important;
}

/* ── Labels ── */
.stTextInput label, .stTextArea label, .stSelectbox label,
.stNumberInput label, .stCheckbox label, .stFileUploader label {
    font-family: 'Space Mono', monospace !important;
    color: #888888 !important;
    font-size: 0.7rem !important;
    font-weight: 400 !important;
    letter-spacing: 0.1em !important;
    text-transform: uppercase !important;
}

/* ── Containers ── */
[data-testid="stVerticalBlock"] > div > div[data-testid="stVerticalBlockBorderWrapper"] {
    background: #ffffff !important;
    border: 2px solid #111111 !important;
    border-radius: 0px !important;
    padding: 1.4rem !important;
}

/* ── Métricas ── */
[data-testid="metric-container"] {
    background: #111111 !important;
    border: none !important;
    border-radius: 0px !important;
    padding: 1.2rem !important;
}
[data-testid="metric-container"] label {
    font-family: 'Space Mono', monospace !important;
    color: #666666 !important;
    font-size: 0.68rem !important;
    text-transform: uppercase !important;
    letter-spacing: 0.1em !important;
}
[data-testid="stMetricValue"] {
    color: #ff3c00 !important;
    font-weight: 700 !important;
    font-size: 2rem !important;
    letter-spacing: -0.03em;
}

/* ── Tabs ── */
.stTabs [data-baseweb="tab-list"] {
    background: transparent !important;
    border-bottom: 2px solid #111111 !important;
    gap: 0 !important;
}
.stTabs [data-baseweb="tab"] {
    font-family: 'Space Mono', monospace !important;
    color: #aaaaaa !important;
    font-size: 0.72rem !important;
    font-weight: 400 !important;
    letter-spacing: 0.08em !important;
    text-transform: uppercase !important;
    padding: 0.65rem 1.2rem !important;
    border-bottom: 3px solid transparent !important;
    background: transparent !important;
}
.stTabs [aria-selected="true"] {
    color: #111111 !important;
    border-bottom: 3px solid #ff3c00 !important;
    background: transparent !important;
}

/* ── Dataframes ── */
[data-testid="stDataFrame"] {
    border: 2px solid #111111 !important;
    border-radius: 0px !important;
    overflow: hidden;
}

/* ── Divider ── */
hr {
    border: none !important;
    border-top: 2px solid #dddddd !important;
    margin: 1.8rem 0 !important;
}

/* ── Caption ── */
.stCaption, [data-testid="stCaptionContainer"] {
    font-family: 'Space Mono', monospace !important;
    color: #999999 !important;
    font-size: 0.72rem !important;
}

/* ── File uploader ── */
[data-testid="stFileUploaderDropzone"] {
    background: #ffffff !important;
    border: 2px dashed #cccccc !important;
    border-radius: 0px !important;
}

/* ── Progress bar ── */
.stProgress > div > div > div {
    background: #ff3c00 !important;
    border-radius: 0px !important;
}

/* ── Expander ── */
[data-testid="stExpander"] {
    background: #ffffff !important;
    border: 2px solid #dddddd !important;
    border-radius: 0px !important;
}
[data-testid="stExpander"] summary {
    font-family: 'Space Mono', monospace !important;
    color: #555555 !important;
    font-size: 0.75rem !important;
    text-transform: uppercase !important;
    letter-spacing: 0.08em !important;
}

/* ── Checkbox ── */
[data-testid="stCheckbox"] span {
    color: #333333 !important;
    font-size: 0.88rem !important;
}

/* ── Alertas ── */
[data-testid="stAlert"] {
    border-radius: 0px !important;
    border-left: 3px solid #ff3c00 !important;
}
</style>
""", unsafe_allow_html=True)

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False

if not st.session_state.logged_in:
    login_form()
    st.stop()

username_actual = st.session_state.username
is_admin = st.session_state.get("is_admin", False)

# ─────────────────────────────────────────────
# NAVEGACIÓN
# ─────────────────────────────────────────────

if "pagina" not in st.session_state:
    st.session_state.pagina = "Inicio"

st.sidebar.title("📧 Prospección en Frío")
st.sidebar.caption(f"👤 {st.session_state.get('user_nombre', username_actual)}")

pagina_sel = st.sidebar.radio("", PAGINAS, index=PAGINAS.index(st.session_state.pagina))
if pagina_sel != st.session_state.pagina:
    st.session_state.pagina = pagina_sel
pagina = st.session_state.pagina

st.sidebar.divider()
if st.sidebar.button("Cerrar sesión"):
    for k in ["logged_in", "username", "user_nombre", "is_admin", "pagina"]:
        st.session_state.pop(k, None)
    st.rerun()


# ═══════════════════════════════════════════════
# INICIO
# ═══════════════════════════════════════════════

if pagina == "Inicio":
    nombre_display = st.session_state.get('user_nombre', username_actual)
    st.markdown(f"""
    <div style="margin-bottom:0.5rem;">
        <h1 style="margin-bottom:0.2rem;">Bienvenido, {nombre_display}</h1>
        <p style="color:#606080; font-size:0.9rem; margin:0;">
            ¿Qué quieres hacer hoy?
        </p>
    </div>
    """, unsafe_allow_html=True)

    rem, _ = get_email_credentials_usuario(username_actual)
    if not rem:
        st.warning("⚠️ Aún no has configurado tu correo. Ve a **Mi configuración → Correo** para poder enviar.")

    col1, col2 = st.columns(2)
    with col1:
        st.markdown("""
        <div style="
            background:#111111; border:none;
            padding:2rem 1.6rem 1.4rem;
        ">
            <div style="font-family:'Space Mono',monospace; font-size:0.65rem;
                        letter-spacing:0.14em; text-transform:uppercase;
                        color:#ff3c00; margin-bottom:0.8rem;">
                01 — Campaña
            </div>
            <div style="font-family:'Space Grotesk',sans-serif; font-weight:700;
                        font-size:1.4rem; color:#f2f0eb; margin-bottom:0.6rem;
                        text-transform:uppercase; letter-spacing:-0.02em;">
                Envío masivo
            </div>
            <div style="color:#666666; font-size:0.84rem; line-height:1.6;">
                Sube un Excel con tus prospectos y envía correos personalizados a todos de un solo clic.
            </div>
        </div>
        """, unsafe_allow_html=True)
        st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)
        if st.button("Ir a Envío masivo →", type="primary", key="btn_masivo"):
            st.session_state.pagina = "Envío masivo"
            st.rerun()
    with col2:
        st.markdown("""
        <div style="
            background:#f2f0eb; border:2px solid #111111;
            padding:2rem 1.6rem 1.4rem;
        ">
            <div style="font-family:'Space Mono',monospace; font-size:0.65rem;
                        letter-spacing:0.14em; text-transform:uppercase;
                        color:#888888; margin-bottom:0.8rem;">
                02 — Prospecto
            </div>
            <div style="font-family:'Space Grotesk',sans-serif; font-weight:700;
                        font-size:1.4rem; color:#111111; margin-bottom:0.6rem;
                        text-transform:uppercase; letter-spacing:-0.02em;">
                Envío individual
            </div>
            <div style="color:#666666; font-size:0.84rem; line-height:1.6;">
                Envía un correo personalizado a un prospecto específico con la plantilla activa.
            </div>
        </div>
        """, unsafe_allow_html=True)
        st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)
        if st.button("Ir a Envío individual →", type="primary", key="btn_individual"):
            st.session_state.pagina = "Envío individual"
            st.rerun()
    

    st.divider()
    campanas_df = get_campanas(solo_usuario=None if is_admin else username_actual)
    if not campanas_df.empty:
        total_enviados = int(campanas_df['total'].sum())
        total_exitosos = int(campanas_df['exitosos'].sum())
        tasa = round(total_exitosos / total_enviados * 100, 1) if total_enviados else 0
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Campañas", len(campanas_df))
        m2.metric("Correos enviados", total_enviados)
        m3.metric("Exitosos", total_exitosos)
        m4.metric("Tasa de éxito", f"{tasa}%")


# ═══════════════════════════════════════════════
# ENVÍO MASIVO
# ═══════════════════════════════════════════════

elif pagina == "Envío masivo":
    st.title("Envío masivo")

    columnas = get_columnas_excel()
    col_correo = get_config("columna_correo", "Correo")
    cols_sin_correo = [c for c in columnas if c != col_correo]
    asunto_plantilla, cuerpo_plantilla = get_plantilla_usuario(username_actual)

    rem, _ = get_email_credentials_usuario(username_actual)
    if not rem:
        st.error("Sin credenciales de correo. Ve a **Mi configuración → Correo** antes de enviar.")
        st.stop()

    st.info(f"El Excel debe tener las columnas: **{', '.join(columnas)}**. Separa con `;` si hay varios destinatarios.")

    st.subheader("1. Nombre de la campaña")
    nombre_campana = st.text_input("Nombre", placeholder="Ej: Prospección Julio 2026 — PYMES CDMX")

    st.subheader("2. Excel de prospectos")
    col_desc, col_btn = st.columns([3, 1])
    with col_desc:
        st.caption(f"Columnas requeridas: **{', '.join(columnas)}**")
    with col_btn:
        st.download_button(
            "⬇️ Plantilla Excel",
            generar_plantilla_excel(columnas),
            file_name="plantilla_prospectos.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    archivo_excel = st.file_uploader("Sube tu Excel lleno", type=["xlsx"])
    df_prospectos = None
    if archivo_excel:
        try:
            df_prospectos = pd.read_excel(archivo_excel, engine="openpyxl")
            df_prospectos.columns = [str(c).strip() for c in df_prospectos.columns]
            df_prospectos = df_prospectos.dropna(how="all").reset_index(drop=True)
            faltantes = [c for c in columnas if c not in df_prospectos.columns]
            if faltantes:
                st.error(f"Al Excel le faltan: **{', '.join(faltantes)}**")
                df_prospectos = None
            else:
                st.success(f"✅ {len(df_prospectos)} prospecto(s) cargados.")
                if len(df_prospectos) > 400:
                    st.warning("⚠️ Gmail permite ~500 correos diarios. Considera dividirlo en varias campañas.")
                st.dataframe(df_prospectos[columnas].head(5))
        except Exception as e:
            st.error(f"No pude leer el archivo: {e}")

    st.subheader("3. Asunto y cuerpo")
    st.caption(f"Placeholders: {' · '.join(f'**({c})**' for c in cols_sin_correo)}")
    asunto = st.text_input("Asunto", value=asunto_plantilla)
    cuerpo = st.text_area("Cuerpo", value=cuerpo_plantilla, height=250)

    invalidos = validar_placeholders(asunto, cuerpo, columnas)
    if invalidos:
        st.warning(f"Placeholders no reconocidos: {', '.join(f'({p})' for p in invalidos)}")

    st.subheader("4. Adjuntos (opcional)")
    adjuntos_masivo = st.file_uploader("Archivos adjuntos", type=None, key="adj_masivo", accept_multiple_files=True)

    if df_prospectos is not None and nombre_campana.strip() and not invalidos:
        st.subheader("5. Vista previa y envío")
        primera_fila = df_prospectos.iloc[0].to_dict()
        asunto_prev = aplicar_placeholders(asunto, primera_fila, columnas)
        cuerpo_prev = aplicar_placeholders(cuerpo, primera_fila, columnas)
        nombre_prev = formatear_multi(str(primera_fila.get(cols_sin_correo[0], ""))) if cols_sin_correo else ""

        with st.expander(f"👁 Preview — {nombre_prev}", expanded=True):
            st.markdown(f"**Asunto:** {asunto_prev}")
            st.divider()
            st.text(cuerpo_prev)

        if st.button(f"Enviar a {len(df_prospectos)} prospecto(s)", type="primary"):
            lista_adj = [(a.read(), a.name) for a in adjuntos_masivo] if adjuntos_masivo else []
            adj_nombre_str = ", ".join(a.name for a in adjuntos_masivo) if adjuntos_masivo else None
            campana_id = crear_campana(nombre_campana.strip(), asunto, cuerpo, "masivo",
                                       username_actual, adjunto_nombre=adj_nombre_str)
            barra = st.progress(0, text="Iniciando envío...")
            enviados, errores = 0, 0
            for i, (_, fila) in enumerate(df_prospectos.iterrows()):
                fila_dict = fila.to_dict()
                correos_str = str(fila_dict.get(col_correo, ""))
                correos_lista = [c.strip() for c in correos_str.split(';') if c.strip()]
                if not correos_lista:
                    registrar_correo(campana_id, "", correos_str, "", "", "error",
                                     fila_datos=fila_dict, error="Sin correo")
                    errores += 1
                    continue
                asunto_f = aplicar_placeholders(asunto, fila_dict, columnas)
                cuerpo_f = aplicar_placeholders(cuerpo, fila_dict, columnas)
                nombres_f = formatear_multi(str(fila_dict.get(cols_sin_correo[0], ""))) if cols_sin_correo else correos_str
                ok, msg = enviar_correo(correos_lista, asunto_f, cuerpo_f, adjuntos=lista_adj, username=username_actual)
                registrar_correo(campana_id, nombres_f, correos_str, asunto_f, cuerpo_f,
                                 "enviado" if ok else "error", fila_datos=fila_dict, error=None if ok else msg)
                enviados += (1 if ok else 0)
                errores += (0 if ok else 1)
                barra.progress((i + 1) / len(df_prospectos), text=f"Enviando {i+1}/{len(df_prospectos)}…")
                time.sleep(0.3)
            barra.empty()
            if errores == 0:
                st.success(f"✅ {enviados} correo(s) enviados.")
            else:
                st.warning(f"✅ {enviados} enviados · ❌ {errores} con error.")
    elif df_prospectos is not None and not nombre_campana.strip():
        st.info("Escribe un nombre para la campaña antes de continuar.")


# ═══════════════════════════════════════════════
# ENVÍO INDIVIDUAL
# ═══════════════════════════════════════════════

elif pagina == "Envío individual":
    st.title("Envío individual")

    columnas = get_columnas_excel()
    col_correo = get_config("columna_correo", "Correo")
    cols_sin_correo = [c for c in columnas if c != col_correo]
    asunto_plantilla, cuerpo_plantilla = get_plantilla_usuario(username_actual)

    rem, _ = get_email_credentials_usuario(username_actual)
    if not rem:
        st.error("Sin credenciales de correo. Ve a **Mi configuración → Correo** antes de enviar.")
        st.stop()

    st.caption(f"Placeholders: {' · '.join(f'**({c})**' for c in cols_sin_correo)}")

    st.subheader("Datos del prospecto")
    campos = {}
    n_cols = min(len(columnas), 3)
    col_inputs = st.columns(n_cols)
    for idx, col in enumerate(columnas):
        with col_inputs[idx % n_cols]:
            hint = f"{col} (separa con ; si son varios)" if col == col_correo else col
            campos[col] = st.text_input(hint, key=f"ind_{col}")

    st.subheader("Asunto y cuerpo")
    asunto_ind = st.text_input("Asunto", value=asunto_plantilla, key="asunto_ind")
    cuerpo_ind = st.text_area("Cuerpo", value=cuerpo_plantilla, height=250, key="cuerpo_ind")
    adjuntos_ind = st.file_uploader("Adjunto(s) (opcional)", key="adj_ind", accept_multiple_files=True)

    correo_ingresado = campos.get(col_correo, "").strip()
    if correo_ingresado:
        asunto_prev = aplicar_placeholders(asunto_ind, campos, columnas)
        cuerpo_prev = aplicar_placeholders(cuerpo_ind, campos, columnas)
        nombre_prev = formatear_multi(campos.get(cols_sin_correo[0], "")) if cols_sin_correo else correo_ingresado

        with st.expander(f"👁 Preview — {nombre_prev}", expanded=True):
            st.markdown(f"**Asunto:** {asunto_prev}")
            st.divider()
            st.text(cuerpo_prev)

        if st.button("Enviar correo", type="primary"):
            correos_lista = [c.strip() for c in correo_ingresado.split(';') if c.strip()]
            lista_adj = [(a.read(), a.name) for a in adjuntos_ind] if adjuntos_ind else []
            adj_nombre_str = ", ".join(a.name for a in adjuntos_ind) if adjuntos_ind else None
            campana_id = crear_campana(f"Individual — {nombre_prev} — {ahora_cdmx()[:10]}",
                                       asunto_ind, cuerpo_ind, "individual", username_actual,
                                       adjunto_nombre=adj_nombre_str)
            ok, msg = enviar_correo(correos_lista, asunto_prev, cuerpo_prev, adjuntos=lista_adj,
                                    username=username_actual)
            registrar_correo(campana_id, nombre_prev, correo_ingresado, asunto_prev, cuerpo_prev,
                             "enviado" if ok else "error", fila_datos=campos, error=None if ok else msg)
            if ok:
                st.success(f"✅ Correo enviado a {', '.join(correos_lista)}.")
            else:
                st.error(f"❌ Error: {msg}")
    else:
        st.info(f"Llena el campo **{col_correo}** para ver la vista previa.")


# ═══════════════════════════════════════════════
# DASHBOARD
# ═══════════════════════════════════════════════

elif pagina == "Dashboard":
    st.title("Dashboard")

    campanas_df = get_campanas(solo_usuario=None if is_admin else username_actual)

    if campanas_df.empty:
        st.info("No hay campañas registradas todavía.")
        st.stop()

    total_enviados = int(campanas_df['total'].sum())
    total_exitosos = int(campanas_df['exitosos'].sum())
    tasa = (total_exitosos / total_enviados * 100) if total_enviados else 0

    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("Campañas", len(campanas_df))
    m2.metric("Correos enviados", total_enviados)
    m3.metric("Exitosos", total_exitosos)
    m4.metric("Con error", int(campanas_df['errores'].sum()))
    m5.metric("Tasa de éxito", f"{tasa:.1f}%")

    conn = get_conn()
    where_dia = f"WHERE c.usuario='{username_actual}'" if not is_admin else ""
    df_por_dia = pd.read_sql_query(
        f"""SELECT DATE(ce.enviado_en) AS dia,
           COUNT(*) AS total,
           SUM(CASE WHEN ce.estado='enviado' THEN 1 ELSE 0 END) AS exitosos
           FROM correos_enviados ce
           JOIN campanas c ON ce.campana_id = c.id
           {where_dia}
           GROUP BY dia ORDER BY dia""", conn)
    conn.close()

    if not df_por_dia.empty:
        st.subheader("Envíos por día")
        fig = px.bar(df_por_dia, x="dia", y="exitosos",
                     labels={"dia": "Día", "exitosos": "Correos enviados"},
                     color_discrete_sequence=["#ff3c00"])
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
        st.plotly_chart(fig)

    st.subheader("Campañas")
    cols_mostrar = ["nombre","tipo","usuario","total","exitosos","errores","creado_en"] if is_admin \
        else ["nombre","tipo","total","exitosos","errores","creado_en"]
    st.dataframe(campanas_df[cols_mostrar].rename(columns={
        "nombre":"Campaña","tipo":"Tipo","usuario":"Usuario",
        "total":"Total","exitosos":"✅","errores":"❌","creado_en":"Fecha"}))

    st.subheader("Detalle por campaña")

    if is_admin:
        filtro_u_dash = st.selectbox("Ver campañas de:", ["Todos"] + get_usuarios()["username"].tolist(), key="dash_filtro_u")
        if filtro_u_dash != "Todos":
            campanas_df = campanas_df[campanas_df["usuario"] == filtro_u_dash]

    campana_id_sel = st.selectbox(
        "Campaña",
        campanas_df["id"].tolist(),
        format_func=lambda x: campanas_df.loc[campanas_df["id"] == x, "nombre"].iloc[0]
    )
    df_correos = get_correos_campana(campana_id_sel)
    if not df_correos.empty:
        respondidos = int(df_correos.get("respondido", pd.Series([0])).fillna(0).sum()) \
            if "respondido" in df_correos.columns else 0
        exitosos_total = int((df_correos["estado"] == "enviado").sum())
        c1, c2, c3 = st.columns(3)
        c1.metric("Enviados", exitosos_total)
        c2.metric("Ya respondieron ✅", respondidos)
        c3.metric("Sin respuesta aún", exitosos_total - respondidos)

        st.caption("Marca a quién ya respondió — esos se excluirán automáticamente de futuros seguimientos.")
        for _, row in df_correos[df_correos["estado"] == "enviado"].iterrows():
            ya_respondio = bool(row.get("respondido", 0))
            col_nombre, col_correo_col, col_seg, col_btn = st.columns([3, 3, 1, 2])
            col_nombre.write(row["nombres"])
            col_correo_col.write(row["correos"])
            n_seg = len(get_seguimientos_por_correo(int(row["id"])))
            col_seg.caption(f"{n_seg} seg.")
            lbl = "✅ Respondió" if ya_respondio else "📭 Sin respuesta"
            if col_btn.button(lbl, key=f"resp_{row['id']}",
                              type="secondary" if ya_respondio else "primary"):
                marcar_respondido(int(row["id"]), not ya_respondio)
                st.rerun()

        errores_df = df_correos[df_correos["estado"] == "error"]
        if not errores_df.empty:
            with st.expander(f"Ver {len(errores_df)} correo(s) con error"):
                st.dataframe(errores_df[["nombres","correos","error"]].rename(columns={
                    "nombres":"Nombre","correos":"Correo(s)","error":"Error"}))
    else:
        st.info("Esta campaña no tiene correos registrados.")


# ═══════════════════════════════════════════════
# SEGUIMIENTOS
# ═══════════════════════════════════════════════

elif pagina == "Seguimientos":
    st.title("Seguimientos")

    solo_usuario = None if is_admin else username_actual
    campanas_df = get_campanas(solo_usuario=solo_usuario)
    campanas_con_enviados = campanas_df[campanas_df["exitosos"] > 0] if not campanas_df.empty else campanas_df

    if campanas_con_enviados.empty:
        st.info("No hay campañas con correos enviados todavía.")
        st.stop()

    if is_admin:
        filtro_u = st.selectbox("Campañas de:", ["Todos"] + get_usuarios()["username"].tolist(), key="seg_filtro_u")
        if filtro_u != "Todos":
            campanas_con_enviados = campanas_con_enviados[campanas_con_enviados["usuario"] == filtro_u]

    campana_id_sel = st.selectbox(
        "Campaña de origen",
        campanas_con_enviados["id"].tolist(),
        format_func=lambda x: campanas_con_enviados.loc[campanas_con_enviados["id"] == x, "nombre"].iloc[0]
    )

    usuario_campana = campanas_con_enviados.loc[campanas_con_enviados["id"] == campana_id_sel, "usuario"].iloc[0]

    df_exitosos = get_correos_campana(campana_id_sel, solo_exitosos=True)
    if df_exitosos.empty:
        st.info("No hay correos exitosos en esta campaña.")
        st.stop()

    df_exitosos = df_exitosos.copy()
    df_exitosos["n_seg"] = df_exitosos["id"].apply(lambda i: len(get_seguimientos_por_correo(int(i))))
    if "respondido" not in df_exitosos.columns:
        df_exitosos["respondido"] = 0
    df_exitosos["respondido"] = df_exitosos["respondido"].fillna(0).astype(int)

    df_sin_respuesta = df_exitosos[df_exitosos["respondido"] == 0]
    df_respondidos   = df_exitosos[df_exitosos["respondido"] == 1]

    if not df_respondidos.empty:
        st.info(f"✅ **{len(df_respondidos)}** contacto(s) marcado(s) como 'Ya respondió' — excluidos automáticamente.")

    incluir_respondidos = False
    if not df_respondidos.empty:
        incluir_respondidos = st.checkbox(
            f"Incluir también a los {len(df_respondidos)} que ya respondieron",
            value=False, key="chk_incluir_respondidos"
        )

    df_para_seguimiento = df_exitosos if incluir_respondidos else df_sin_respuesta

    if df_para_seguimiento.empty:
        st.warning("Todos los contactos ya respondieron.")
        st.stop()

    seleccionar_todos = st.checkbox("Seleccionar todos", value=True, key="chk_sel_todos_seg")
    if seleccionar_todos:
        ids_sel = df_para_seguimiento["id"].tolist()
    else:
        def label_seg(x):
            row = df_para_seguimiento[df_para_seguimiento["id"] == x].iloc[0]
            partes = [f"{row['nombres']} — {row['correos']}"]
            if row["n_seg"] > 0:
                partes.append(f"{row['n_seg']} seg.")
            if row["respondido"]:
                partes.append("✅ ya respondió")
            return "  ·  ".join(partes)

        ids_sel = st.multiselect(
            "Destinatarios",
            options=df_para_seguimiento["id"].tolist(),
            default=df_para_seguimiento["id"].tolist(),
            format_func=label_seg
        )

    st.caption(f"{len(ids_sel)} destinatario(s) seleccionado(s).")

    columnas = get_columnas_excel()
    col_correo = get_config("columna_correo", "Correo")
    cols_sin_correo = [c for c in columnas if c != col_correo]

    st.subheader("Correo de seguimiento")
    st.caption(f"Placeholders: {' · '.join(f'**({c})**' for c in cols_sin_correo)}")
    asunto_seg = st.text_input("Asunto", placeholder="Seguimiento a mi propuesta para (Empresa)", key="asunto_seg")
    cuerpo_seg = st.text_area("Cuerpo", height=200, key="cuerpo_seg",
                              placeholder="Hola (Nombre),\n\nQuería hacer seguimiento...")
    adjuntos_seg = st.file_uploader("Adjunto(s) (opcional)", key="adj_seg", accept_multiple_files=True)

    invalidos_seg = validar_placeholders(asunto_seg, cuerpo_seg, columnas) if asunto_seg and cuerpo_seg else []
    if invalidos_seg:
        st.warning(f"Placeholders no reconocidos: {', '.join(f'({p})' for p in invalidos_seg)}")

    if ids_sel and asunto_seg and cuerpo_seg and not invalidos_seg:
        if st.button(f"Enviar seguimiento a {len(ids_sel)} destinatario(s)", type="primary"):
            lista_adj = [(a.read(), a.name) for a in adjuntos_seg] if adjuntos_seg else []
            barra = st.progress(0)
            enviados, errores = 0, 0
            for i, id_correo in enumerate(ids_sel):
                fila_correo = df_para_seguimiento[df_para_seguimiento["id"] == id_correo].iloc[0]
                fila_datos = {}
                raw = fila_correo.get("fila_datos")
                if isinstance(raw, str) and raw.strip():
                    try:
                        fila_datos = json.loads(raw)
                    except Exception:
                        fila_datos = {}
                if not fila_datos:
                    fila_datos = {"Nombre": fila_correo["nombres"], col_correo: fila_correo["correos"]}
                asunto_f = aplicar_placeholders(asunto_seg, fila_datos, columnas)
                cuerpo_f = aplicar_placeholders(cuerpo_seg, fila_datos, columnas)
                correos_lista = [c.strip() for c in fila_correo["correos"].split(';') if c.strip()]
                ok, msg = enviar_correo(correos_lista, asunto_f, cuerpo_f, adjuntos=lista_adj, username=usuario_campana)
                registrar_seguimiento(int(id_correo), asunto_f, cuerpo_f, "enviado" if ok else "error", None if ok else msg)
                enviados += (1 if ok else 0)
                errores += (0 if ok else 1)
                barra.progress((i + 1) / len(ids_sel))
                time.sleep(0.3)
            barra.empty()
            if errores == 0:
                st.success(f"✅ {enviados} seguimiento(s) enviados.")
            else:
                st.warning(f"✅ {enviados} enviados · ❌ {errores} con error.")


# ═══════════════════════════════════════════════
# MI CONFIGURACIÓN
# ═══════════════════════════════════════════════

elif pagina == "Mi configuración":
    st.title("Mi configuración")
    st.caption("Esta configuración aplica solo a tu cuenta.")

    tab_correo, tab_plantilla = st.tabs(["Correo", "Plantilla"])

    with tab_correo:
        st.subheader("Mi correo de envío")
        st.caption(
            "Tus correos salen desde aquí. Usa Gmail con verificación en 2 pasos y genera una "
            "**Contraseña de aplicación** en myaccount.google.com/apppasswords."
        )
        row = get_usuario_row(username_actual)
        mi_remitente = st.text_input("Correo remitente", value=row.get("email_remitente", ""))
        mi_password = st.text_input("Contraseña de aplicación", value=row.get("email_password", ""), type="password")
        if st.button("Guardar mi correo", type="primary", key="btn_guardar_mi_correo"):
            set_usuario_email(username_actual, mi_remitente, mi_password)
            st.success("Correo guardado.")

    with tab_plantilla:
        st.subheader("Mi plantilla de correo")
        st.caption("Esta plantilla se carga automáticamente en tus envíos.")
        columnas = get_columnas_excel()
        col_correo = get_config("columna_correo", "Correo")
        cols_sin_correo = [c for c in columnas if c != col_correo]
        st.caption(f"Placeholders: {' · '.join(f'**({c})**' for c in cols_sin_correo)}")
        mi_asunto_actual, mi_cuerpo_actual = get_plantilla_usuario(username_actual)
        mi_asunto = st.text_input("Asunto", value=mi_asunto_actual, key="mi_asunto_plantilla")
        mi_cuerpo = st.text_area("Cuerpo", value=mi_cuerpo_actual, height=300, key="mi_cuerpo_plantilla")
        if st.button("Guardar mi plantilla", type="primary", key="btn_guardar_mi_plantilla"):
            set_usuario_plantilla(username_actual, mi_asunto, mi_cuerpo)
            st.success("Plantilla guardada.")

    st.divider()
    st.subheader("Cambiar mi contraseña")
    pwd_actual = st.text_input("Contraseña actual", type="password", key="pwd_actual")
    pwd_nueva = st.text_input("Nueva contraseña", type="password", key="pwd_nueva")
    pwd_confirmar = st.text_input("Confirmar nueva contraseña", type="password", key="pwd_confirmar")
    if st.button("Cambiar contraseña", key="btn_cambiar_pwd_propia"):
        if not verificar_usuario(username_actual, pwd_actual):
            st.error("La contraseña actual no es correcta.")
        elif pwd_nueva != pwd_confirmar:
            st.error("Las contraseñas nuevas no coinciden.")
        elif not pwd_nueva.strip():
            st.error("La nueva contraseña no puede estar vacía.")
        else:
            row = get_usuario_row(username_actual)
            cambiar_password_usuario(int(row["id"]), pwd_nueva)
            st.success("Contraseña actualizada.")


# ═══════════════════════════════════════════════
# ADMINISTRACIÓN
# ═══════════════════════════════════════════════

elif pagina == "Administración":
    verificar_admin()
    st.title("Administración")

    tab_usuarios, tab_columnas, tab_limpiar, tab_seguridad, tab_respaldo = st.tabs([
        "Usuarios", "Columnas del Excel", "Limpiar datos", "Seguridad", "Respaldo"
    ])

    # ── Usuarios ──
    with tab_usuarios:
        st.subheader("Usuarios")
        usuarios_df = get_usuarios()
        st.dataframe(usuarios_df[["username","nombre","activo","email_remitente","creado_en"]].rename(columns={
            "username":"Usuario","nombre":"Nombre","activo":"Activo",
            "email_remitente":"Correo configurado","creado_en":"Creado"}))

        st.markdown("**Crear nuevo usuario**")
        c1, c2, c3 = st.columns(3)
        with c1: new_username = st.text_input("Usuario", key="new_username")
        with c2: new_nombre = st.text_input("Nombre completo", key="new_nombre")
        with c3: new_password = st.text_input("Contraseña inicial", type="password", key="new_password")
        if st.button("Crear usuario", key="btn_crear_usuario"):
            if new_username.strip() and new_password.strip():
                ok, msg = crear_usuario(new_username, new_password, new_nombre)
                if ok:
                    st.success(msg)
                    st.rerun()
                else:
                    st.error(msg)
            else:
                st.error("Usuario y contraseña son obligatorios.")

        st.markdown("**Cambiar contraseña**")
        c1, c2, c3 = st.columns([2, 2, 1])
        with c1: user_pwd = st.selectbox("Usuario", usuarios_df["username"].tolist(), key="sel_user_pwd")
        with c2: nueva_pwd_u = st.text_input("Nueva contraseña", type="password", key="nueva_pwd_u")
        with c3:
            st.write("")
            if st.button("Cambiar", key="btn_cambiar_pwd_admin"):
                if nueva_pwd_u.strip():
                    uid = int(usuarios_df.loc[usuarios_df["username"] == user_pwd, "id"].iloc[0])
                    cambiar_password_usuario(uid, nueva_pwd_u)
                    st.success("Contraseña actualizada.")
                else:
                    st.error("Escribe la nueva contraseña.")

        st.markdown("**Activar / Desactivar usuario**")
        c1, c2 = st.columns([3, 1])
        with c1: user_toggle = st.selectbox("Usuario", usuarios_df["username"].tolist(), key="sel_user_toggle")
        with c2:
            activo_actual = bool(usuarios_df.loc[usuarios_df["username"] == user_toggle, "activo"].iloc[0])
            st.write("")
            if st.button("Desactivar" if activo_actual else "Activar", key="btn_toggle_user"):
                uid = int(usuarios_df.loc[usuarios_df["username"] == user_toggle, "id"].iloc[0])
                toggle_usuario_activo(uid, not activo_actual)
                st.success(f"Usuario {'desactivado' if activo_actual else 'activado'}.")
                st.rerun()

    # ── Columnas del Excel ──
    with tab_columnas:
        st.subheader("Columnas del Excel (globales para todos los usuarios)")
        columnas = get_columnas_excel()
        col_correo = get_config("columna_correo", "Correo")
        st.dataframe(pd.DataFrame({
            "Columna": columnas,
            "Es columna de correo": [c == col_correo for c in columnas]
        }))
        c1, c2 = st.columns([3, 1])
        with c1: nueva_col = st.text_input("Nombre", key="nueva_col_excel")
        with c2:
            st.write("")
            if st.button("Agregar", key="btn_agregar_col"):
                nc = nueva_col.strip()
                if nc and nc not in columnas:
                    columnas.append(nc)
                    set_config("columnas_excel", json.dumps(columnas))
                    st.success("Columna agregada.")
                    st.rerun()
                elif nc in columnas:
                    st.error("Ya existe.")
        if len(columnas) > 1:
            c1, c2 = st.columns([3, 1])
            with c1: col_quitar = st.selectbox("Columna a quitar", columnas, key="col_quitar")
            with c2:
                st.write("")
                if st.button("Quitar", type="secondary", key="btn_quitar_col"):
                    if col_quitar == col_correo:
                        st.error("No puedes quitar la columna de correo.")
                    else:
                        columnas.remove(col_quitar)
                        set_config("columnas_excel", json.dumps(columnas))
                        st.success("Eliminada.")
                        st.rerun()
        c1, c2 = st.columns([3, 1])
        with c1:
            nueva_col_correo = st.selectbox(
                "Columna de correo", columnas,
                index=columnas.index(col_correo) if col_correo in columnas else 0,
                key="sel_col_correo"
            )
        with c2:
            st.write("")
            if st.button("Guardar", key="btn_col_correo"):
                set_config("columna_correo", nueva_col_correo)
                st.success(f"'{nueva_col_correo}' es ahora la columna de correo.")
                st.rerun()

    # ── Limpiar datos ──
    with tab_limpiar:
        st.subheader("Borrar campañas o correos")
        campanas_todas = get_campanas()
        if campanas_todas.empty:
            st.info("No hay campañas registradas.")
        else:
            campana_borrar_id = st.selectbox(
                "Campaña a borrar",
                campanas_todas["id"].tolist(),
                format_func=lambda x: (
                    f"[{campanas_todas.loc[campanas_todas['id']==x,'usuario'].iloc[0]}] "
                    f"{campanas_todas.loc[campanas_todas['id']==x,'nombre'].iloc[0]} "
                    f"({int(campanas_todas.loc[campanas_todas['id']==x,'total'].iloc[0])} correos)"
                ),
                key="sel_campana_borrar"
            )
            nombre_sel = campanas_todas.loc[campanas_todas["id"] == campana_borrar_id, "nombre"].iloc[0]
            st.warning(f"Esto borrará **{nombre_sel}** y todos sus correos. No se puede deshacer.")
            if st.checkbox(f"Confirmo borrar '{nombre_sel}'", key="chk_borrar_campana"):
                if st.button("🗑️ Borrar campaña", type="secondary", key="btn_borrar_campana"):
                    borrar_campana(int(campana_borrar_id))
                    st.success("Campaña eliminada.")
                    st.rerun()

    # ── Seguridad ──
    with tab_seguridad:
        st.subheader("Contraseña de Administración")
        nueva_admin_pwd = st.text_input("Nueva contraseña", type="password")
        if st.button("Actualizar", key="btn_admin_pwd"):
            if nueva_admin_pwd.strip():
                set_config("admin_password", hash_pw(nueva_admin_pwd.strip()))
                st.success("Contraseña actualizada.")
            else:
                st.error("No puede estar vacía.")
        st.divider()
        if st.button("Cerrar sesión de administrador", key="btn_cerrar_admin"):
            st.session_state.is_admin = False
            st.rerun()

    # ── Respaldo ──
    with tab_respaldo:
        st.subheader("Respaldo de la base de datos")
        try:
            conn_r = get_conn()
            tablas_r = {}
            for tabla in ["configuracion", "usuarios", "campanas", "correos_enviados", "seguimientos"]:
                df_t = pd.read_sql_query(f"SELECT * FROM {tabla}", conn_r)
                tablas_r[tabla] = df_t.to_dict("records")
            conn_r.close()
            respaldo_bytes = json.dumps(tablas_r, ensure_ascii=False, default=str).encode("utf-8")
            st.download_button(
                "⬇️ Descargar respaldo (.json)",
                respaldo_bytes,
                file_name=f"respaldo_prospeccion_{ahora_cdmx()[:10]}.json",
                mime="application/json"
            )
        except Exception as e:
            st.error(f"Error al generar respaldo: {e}")
